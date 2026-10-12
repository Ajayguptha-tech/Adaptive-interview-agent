"""Resume text extraction service for PDF and DOCX files."""
import re
import io
import pymupdf as fitz
import docx
from typing import Tuple, Dict, Any


def normalize_text(text: str) -> str:
    """Normalize whitespace and clean extraction artifacts without destroying paragraph structure."""
    if not text:
        return ""
    # Normalize fancy punctuation / bullets
    text = text.replace("\u2022", " • ").replace("\u2023", " • ").replace("\u25e6", " • ")
    text = text.replace("\u2013", "-").replace("\u2014", "-")
    text = text.replace("\u201c", '"').replace("\u201d", '"')
    text = text.replace("\u2018", "'").replace("\u2019", "'")
    text = text.replace("\xa0", " ")  # non-breaking space
    text = text.replace("\t", "    ")

    # Replace multiple spaces with a single space per line
    lines = []
    for line in text.splitlines():
        cleaned_line = re.sub(r"[ ]{2,}", " ", line.strip())
        lines.append(cleaned_line)

    # Collapse multiple consecutive empty lines to at most one empty line
    collapsed = []
    empty_count = 0
    for line in lines:
        if not line:
            empty_count += 1
            if empty_count <= 1:
                collapsed.append("")
        else:
            empty_count = 0
            collapsed.append(line)

    result = "\n".join(collapsed).strip()
    return re.sub(r"\n{3,}", "\n\n", result)


def extract_from_pdf(file_bytes: bytes) -> Tuple[str, bool, str]:
    """
    Extract text from PDF using PyMuPDF.
    Returns: (extracted_text, is_scanned_pdf, warning_message)
    """
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
    except Exception as e:
        raise ValueError(f"Unable to read PDF file. The file may be corrupt or encrypted: {str(e)}")

    page_texts = []
    total_pages = len(doc)
    
    if total_pages == 0:
        return "", False, "The PDF file contains 0 pages."

    for page_idx in range(total_pages):
        page = doc[page_idx]
        text = page.get_text("text")
        if text.strip():
            page_texts.append(text)

    doc.close()
    full_text = "\n\n".join(page_texts)
    normalized = normalize_text(full_text)

    # Detect if PDF is an image-only scan
    if len(normalized) < 50:
        warning = (
            "Notice: This PDF contains little or no extractable digital text. "
            "It appears to be a scanned image or bitmap. "
            "ResumeIQ extracts digital text directly to keep resource usage lightweight, "
            "and does not perform local OCR. Please upload a PDF with a selectable text layer or a DOCX file."
        )
        return normalized, True, warning

    return normalized, False, ""


def extract_from_docx(file_bytes: bytes) -> Tuple[str, bool, str]:
    """
    Extract text from DOCX using python-docx.
    Returns: (extracted_text, is_scanned, warning_message)
    """
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise ValueError(f"Unable to read DOCX document. File may be corrupted: {str(e)}")

    paragraphs = []
    # Extract body paragraphs
    for p in doc.paragraphs:
        if p.text.strip():
            paragraphs.append(p.text.strip())

    # Extract tables
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                paragraphs.append(" | ".join(row_cells))

    full_text = "\n".join(paragraphs)
    normalized = normalize_text(full_text)

    if not normalized:
        return "", False, "DOCX file contains no readable text."

    return normalized, False, ""


def extract_resume_text(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """Main extraction dispatcher for supported file extensions."""
    lower_name = filename.lower()
    if lower_name.endswith(".pdf"):
        text, is_scanned, warning = extract_from_pdf(file_bytes)
        file_type = "pdf"
    elif lower_name.endswith(".docx") or lower_name.endswith(".doc"):
        if lower_name.endswith(".doc"):
            raise ValueError("Legacy .doc format is not supported. Please save or export your resume as .docx or .pdf.")
        text, is_scanned, warning = extract_from_docx(file_bytes)
        file_type = "docx"
    else:
        raise ValueError("Unsupported file format. Please upload a .pdf or .docx file.")

    word_count = len(text.split()) if text else 0
    char_count = len(text)

    return {
        "text": text,
        "is_scanned": is_scanned,
        "warning": warning,
        "file_type": file_type,
        "word_count": word_count,
        "char_count": char_count,
    }
